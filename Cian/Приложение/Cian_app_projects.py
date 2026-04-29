import tkinter as tk
from tkinter import messagebox
import requests
import requests
import datetime
import time
import pandas as pd
import openpyxl
import random
import re
import os
import sys
import json
import re
import numpy as np
import os
import glob
import datetime
from Developer_dict import name_dict, developer_dict
from area_dictionary.Старая_квартирография.step_4_replacement_to_excel import process_data, load_json_data
import json
from enrich.main2 import enrich_dataframe
from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

proxies = {
    "http": "http://STm87nUFS6:6StepJYs2y@185.42.27.210:10270",
    "https": "http://STm87nUFS6:6StepJYs2y@185.42.27.210:10270"
}

def get_unique_filepath(folder_path: str, filename: str) -> str:
    """
    Если файл существует — добавляет _1, _2, _3 ...
    """
    base, ext = os.path.splitext(filename)
    file_path = os.path.join(folder_path, filename)

    counter = 1
    while os.path.exists(file_path):
        file_path = os.path.join(folder_path, f"{base}_{counter}{ext}")
        counter += 1

    return file_path

def resource_path(relative_path):
    """
    Работает и в обычном Python, и в exe
    """
    try:
        base_path = sys._MEIPASS
    except Exception:
        base_path = os.path.abspath(".")

    return os.path.join(base_path, relative_path)


def load_json(path):
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)

def save_flats_to_excel(flats, project, developer, kvartirografia=True, drop_columns=False, change_dates=False):

    df = pd.DataFrame(flats, columns=['Дата обновления',
                                      'Название проекта',
                                      'На англ',
                                      'Промзона',
                                      'Местоположение',
                                      'Метро',
                                      'Расстояние до метро, км',
                                      'Время до метро, мин',
                                      'МЦК/МЦД/БКЛ',
                                      'Расстояние до МЦК/МЦД, км',
                                      'Время до МЦК/МЦД, мин',
                                      'БКЛ',
                                      'Расстояние до БКЛ, км',
                                      'Время до БКЛ, мин',
                                      'Статус',
                                      'Старт',
                                      'Комментарий',
                                      'Девелопер',
                                      'Округ',
                                      'Район',
                                      'Адрес',
                                      'Эскроу',
                                      'Корпус',
                                      'Конструктив',
                                      'Класс',
                                      'Срок сдачи',
                                      'Старый срок сдачи',
                                      'Стадия строительной готовности',
                                      'Договор',
                                      'Тип помещения',
                                      'Отделка',
                                      'Кол-во комнат',
                                      'Площадь, кв.м',
                                      'Цена кв.м, руб.',
                                      'Цена лота, руб.',
                                      'Скидка,%',
                                      'Цена кв.м со ск, руб.',
                                      'Цена со скидкой, руб.',
                                      'секция',
                                      'этаж',
                                      'номер'])

    df["Корпус"] = (
         df["Корпус"]
         .astype(str)  # приводим всё к строкам
         .str.replace(r'(?i)\bкорпус\b\.?\s*', '', regex=True)  # удаляем "корпус"
         .str.strip()  # убираем лишние пробелы
         .replace(['', '-', 'nan', 'NaN'], '1')  # заменяем пустые строки и текстовые NaN на "1"
     )

    # В столбце Корпус, если номер корпуса идёт в скобках, то удаляем всё за скобками, оставляем только то,
    # что в скобках. Если в скобках есть слово очередь, то ничего не трогаем
    df['Корпус'] = df['Корпус'].astype(str)
    # Выделим текст в скобках
    bracket_content = df['Корпус'].str.extract(r'\((.*?)\)', expand=False)
    # Маска: строка содержит скобки
    has_brackets = df['Корпус'].str.contains(r'\(.*?\)', na=False)
    # Маска: в строке есть слово "очередь" где угодно (внутри или снаружи скобок)
    contains_ochered = df['Корпус'].str.contains(r'очередь', case=False, na=False)
    # Маска: есть скобки, но НЕТ "очередь" вообще
    mask = has_brackets & ~contains_ochered
    # Заменяем только те строки, которые соответствуют маске
    df.loc[mask, 'Корпус'] = bracket_content[mask].str.strip()

    # 1. Удаляем текст 'Жилой дом № ' (всё равно на NaN это не повлияет)
    df['Корпус'] = df['Корпус'].replace('Жилой дом № ', '', regex=True)
    # Заменяем строку 'nan' и пустые строки на np.nan
    df['Корпус'] = df['Корпус'].replace(['nan', r'^\s*$'], np.nan, regex=True)
    # 3. Заполняем NaN единицами
    df['Корпус'] = df['Корпус'].fillna('1')

    df["Корпус"] = df["Корпус"].astype(str).str.replace(",", ".", regex=False)

    def clean_name(name: str) -> str:
        name = name.replace('ЖК ', '')
        name = re.sub(r'[\\:*?"<>|«»]', '', name)
        name = re.sub(r'\s+', ' ', name).strip()
        return name

    developer = clean_name(developer)
    df['Название проекта'] = df['Название проекта'].apply(clean_name)
    df["Название проекта"] = df["Название проекта"].replace(name_dict)
    print(df['Девелопер'].unique())
    df['Девелопер'] = df['Девелопер'].apply(clean_name)
    df["Девелопер"] = df["Девелопер"].replace(developer_dict)
    print(df['Девелопер'].unique())

    projects_dict = load_json(
        resource_path("haracteristik_dictionary/projects.json")
    )

    corpus_dict = load_json(
        resource_path("changing_haracteristik_dictionary/projects.json")
    )

    if kvartirografia:
        area_dict = load_json(
            resource_path("area_dictionary/output.json")
        )
    else:
        area_dict = None

    df = enrich_dataframe(
        df,
        projects_dict,
        corpus_dict,
        area_dict
    )

    if drop_columns:
        df = df.drop(columns=['секция', 'этаж', 'номер'], errors='ignore')
    # --- 1. Приводим цены к числу ПЕРЕД всеми вычислениями ---

    for col in ['Цена лота, руб.', 'Цена со скидкой, руб.']:
        df[col] = (
            df[col]
            .astype(str)
            .str.replace(' ', '', regex=False)
            .str.replace('\xa0', '', regex=False)
            .replace(['', 'None', 'nan'], np.nan)
        )
        df[col] = pd.to_numeric(df[col], errors='coerce')

    # --- 2. Приводим площадь ---
    df['Площадь, кв.м'] = pd.to_numeric(df['Площадь, кв.м'], errors='coerce')

    # --- 3. Цена со скидкой ---
    df['Цена со скидкой, руб.'] = (
        df['Цена со скидкой, руб.']
        .replace(0, np.nan)
        .fillna(df['Цена лота, руб.'])
    )

    # --- 4. Безопасное деление ---
    df['Цена кв.м, руб.'] = np.where(
        df['Площадь, кв.м'] > 0,
        df['Цена лота, руб.'] / df['Площадь, кв.м'],
        np.nan
    )

    df['Цена кв.м со ск, руб.'] = np.where(
        df['Площадь, кв.м'] > 0,
        df['Цена со скидкой, руб.'] / df['Площадь, кв.м'],
        np.nan
    )

    df['Скидка,%'] = np.where(
        df['Цена лота, руб.'] > 0,
        1 - (df['Цена со скидкой, руб.'] / df['Цена лота, руб.']),
        np.nan
    )

    # --- 5. Округление ---
    df['Цена кв.м, руб.'] = df['Цена кв.м, руб.'].round(2)
    df['Цена кв.м со ск, руб.'] = df['Цена кв.м со ск, руб.'].round(2)
    df['Скидка,%'] = df['Скидка,%'].round(2)


    print(df[['Корпус', 'Кол-во комнат', 'Площадь, кв.м', 'Цена лота, руб.', 'Цена со скидкой, руб.']].info())
    print(f'')
    print(f'Число лотов: {len(df)}')
    print(f'')
    print(f'Типы отделки: {df['Отделка'].value_counts()}')
    print(f'')
    print(f'Проекты: {df['Название проекта'].value_counts()}')
    print(f'')


    current_date = datetime.date.today()
    project_root = os.path.dirname(os.path.abspath(__file__))
    base_path = os.path.join(project_root, "1_FILES")
    folder_path = os.path.join(base_path, str(current_date))
    if not os.path.exists(folder_path):
        os.makedirs(folder_path)
    project = re.sub(r'[<>:"/\\|?*]', '_', project)
    filename = f"{developer}_{project}_{current_date}.xlsx"
    file_path = get_unique_filepath(folder_path, filename)
    df.to_excel(file_path, index=False)
    # открываем файл и применяем выравнивание
    wb = load_workbook(file_path)
    ws = wb.active

    center_alignment = Alignment(horizontal='center', vertical='center')

    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = center_alignment

    # Автоподбор ширины колонок
    for col in ws.columns:
        max_length = 0
        col_letter = get_column_letter(col[0].column)  # Получаем буквенное обозначение колонки

        for cell in col:
            try:
                if cell.value:
                    # Учитываем длину текста в ячейке
                    cell_length = len(str(cell.value))
                    if cell_length > max_length:
                        max_length = cell_length
            except:
                pass

        # Устанавливаем ширину колонки (добавляем 2 для отступов)
        adjusted_width = max_length + 2
        ws.column_dimensions[col_letter].width = adjusted_width

    wb.save(file_path)
    print(f"✅ Данные сохранены в файл: {file_path}")



def parse_cian(cian_id):

    print(f'Первоначальный IP: {requests.get("https://ipinfo.io/json", proxies=proxies).json()}')

    decoration_dict = {'preFine' : 'Предчистовая', 'fine' : 'С отделкой', 'without' : 'Без отделки', 'fineWithFurniture' : 'С отделкой и доп опциями'}
    decoration_list = ['preFine', 'fine', 'without', 'fineWithFurniture']


    parsim = [int(cian_id)]

    cookies = {
        '_CIAN_GK': '38928be9-bba1-4562-8d8e-71aa9dfb2ba9',
        '_ym_uid': '174161324651361127',
        'adrcid': 'Ad53EZahiTy4QvZYZHYhh0Q',
        'adrdel': '1744094487237',
        'acs_3': '%7B%22hash%22%3A%221aa3f9523ee6c2690cb34fc702d4143056487c0d%22%2C%22nst%22%3A1744181465976%2C%22sl%22%3A%7B%22224%22%3A1744095065976%2C%221228%22%3A1744095065976%7D%7D',
        'ma_id': '6225667261741613246584',
        '__ai_fp_uuid': '245d903c22bdc927%3A15',
        '_gcl_au': '1.1.9818463.1769411842',
        '_ym_d': '1773209373',
        '_ga': 'GA1.1.1538482319.1774343544',
        'tmr_lvid': '61ae9374a9f1699406db7cc31ef00775',
        'tmr_lvidTS': '1741613242260',
        'transport-accessibility_onboarding_counter': '3',
        'uxs_uid': '92604860-28f8-11f1-a98a-bba19a4d4807',
        'uxfb_usertype': 'searcher',
        'cookie_agreement_accepted': '1',
        'newbuilding-card-desktop-fichering-frontend.builder_chat_onboarding_shown': '1',
        'map_preview_onboarding_counter': '3',
        'login_button_tooltip_key': '1',
        'frontend-serp.header_builder_chat_onboarding_shown': '1',
        'last_paid_utm': '?utm_source=yandex&utm_medium=cpc&utm_content=kw:205657295806|ad:1889238602131455357|grp:5657295806|drf:no|dev:desktop|p:premium|n:2|reg:213|s:none&utm_term=|205657295806&utm_campaign=b2c_nov_mskmo_perf_mix_search_dsa_feed_general_drr_arwm_703812141&etext=2202.h07M9Tlhg3N9CwbhNMoyrVsLWbdrg5CRpbuLNhiEcIVoTY8yPfNz7DVvSxFMf63mtQfzGVZPsUeL_PUMqGvCe3BsenN6dGFxZmZlYWRmaWY.c0ab089525eaca3735ee4aa1b318395f0703704e&yclid=3873403374751711231',
        'sopr_utm': '%7B%22utm_source%22%3A+%22yandex%22%2C+%22utm_medium%22%3A+%22organic%22%7D',
        'newbuilding-search-frontend.builder_chat_onboarding_shown': '1',
        'frontend-serp.chatTooltipAnimationShown': '1',
        'newbuilding-search-frontend.chatTooltipAnimationShown': '1',
        'newbuilding-search-frontend.chatAnimationShownCount': '107',
        'frontend-serp.offer_chat_onboarding_shown': '1',
        'frontend-serp.chatAnimationShownCount': '7',
        'newbuilding-search-frontend.chatAnimationCounter': '111',
        'newbuilding-search-frontend.chatAnimationPrevPath': '%2Fnovostroyki%2F',
        'frontend-offer-card.builder_chat_onboarding_shown': '1',
        'frontend-offer-card.newbuilding_broker_onboarding_shown': '1',
        'frontend-serp.chatAnimationCounter': '10',
        'frontend-serp.chatAnimationPrevPath': '%2Fkupit-kvartiru-novostroyki%2F',
        'countCallNowPopupShowed': '1%3A1775476258941',
        'forever_region_id': '4959',
        'forever_region_name': '%D0%A0%D0%BE%D1%81%D1%82%D0%BE%D0%B2%D1%81%D0%BA%D0%B0%D1%8F%20%D0%BE%D0%B1%D0%BB%D0%B0%D1%81%D1%82%D1%8C',
        'forever_main_town_region_id': '4959',
        'session_region_id': '4593',
        'session_main_town_region_id': '1',
        'login_mro_popup': '1',
        '_ym_isad': '2',
        'newbuilding_mortgage_payment_filter_onboarding': '1',
        'sopr_session': '984fe57857c04e92',
        '_ym_visorc': 'b',
        '_yasc': 'TRKPtQwX100Llb+e03dywP4BMM51qmiqnpIri2vjwENZelIZncfgiKhy9oA6lHkCD45B0dE=',
        '_yasc': '0SMjQFnVkpPLd+wWAfEiWFOgNTsWxvn3eJZjAgkulSdK+kvhtbHsGVU0Gbe2z9s/OyK/5r8=',
        '_ga_3369S417EL': 'GS2.1.s1775561722$o34$g1$t1775562653$j60$l0$h0',
    }

    headers = {
        'accept': '*/*',
        'accept-language': 'ru-RU,ru;q=0.9,en-GB;q=0.8,en;q=0.7,en-US;q=0.6',
        'content-type': 'application/json',
        'origin': 'https://www.cian.ru',
        'priority': 'u=1, i',
        'referer': 'https://www.cian.ru/',
        'sec-ch-ua': '"Chromium";v="146", "Not-A.Brand";v="24", "Google Chrome";v="146"',
        'sec-ch-ua-mobile': '?0',
        'sec-ch-ua-platform': '"Windows"',
        'sec-fetch-dest': 'empty',
        'sec-fetch-mode': 'cors',
        'sec-fetch-site': 'same-site',
        'user-agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/146.0.0.0 Safari/537.36',
        # 'cookie': '_CIAN_GK=38928be9-bba1-4562-8d8e-71aa9dfb2ba9; _ym_uid=174161324651361127; adrcid=Ad53EZahiTy4QvZYZHYhh0Q; adrdel=1744094487237; acs_3=%7B%22hash%22%3A%221aa3f9523ee6c2690cb34fc702d4143056487c0d%22%2C%22nst%22%3A1744181465976%2C%22sl%22%3A%7B%22224%22%3A1744095065976%2C%221228%22%3A1744095065976%7D%7D; ma_id=6225667261741613246584; __ai_fp_uuid=245d903c22bdc927%3A15; _gcl_au=1.1.9818463.1769411842; _ym_d=1773209373; _ga=GA1.1.1538482319.1774343544; tmr_lvid=61ae9374a9f1699406db7cc31ef00775; tmr_lvidTS=1741613242260; transport-accessibility_onboarding_counter=3; uxs_uid=92604860-28f8-11f1-a98a-bba19a4d4807; uxfb_usertype=searcher; cookie_agreement_accepted=1; newbuilding-card-desktop-fichering-frontend.builder_chat_onboarding_shown=1; map_preview_onboarding_counter=3; login_button_tooltip_key=1; frontend-serp.header_builder_chat_onboarding_shown=1; last_paid_utm=?utm_source=yandex&utm_medium=cpc&utm_content=kw:205657295806|ad:1889238602131455357|grp:5657295806|drf:no|dev:desktop|p:premium|n:2|reg:213|s:none&utm_term=|205657295806&utm_campaign=b2c_nov_mskmo_perf_mix_search_dsa_feed_general_drr_arwm_703812141&etext=2202.h07M9Tlhg3N9CwbhNMoyrVsLWbdrg5CRpbuLNhiEcIVoTY8yPfNz7DVvSxFMf63mtQfzGVZPsUeL_PUMqGvCe3BsenN6dGFxZmZlYWRmaWY.c0ab089525eaca3735ee4aa1b318395f0703704e&yclid=3873403374751711231; sopr_utm=%7B%22utm_source%22%3A+%22yandex%22%2C+%22utm_medium%22%3A+%22organic%22%7D; newbuilding-search-frontend.builder_chat_onboarding_shown=1; frontend-serp.chatTooltipAnimationShown=1; newbuilding-search-frontend.chatTooltipAnimationShown=1; newbuilding-search-frontend.chatAnimationShownCount=107; frontend-serp.offer_chat_onboarding_shown=1; frontend-serp.chatAnimationShownCount=7; newbuilding-search-frontend.chatAnimationCounter=111; newbuilding-search-frontend.chatAnimationPrevPath=%2Fnovostroyki%2F; frontend-offer-card.builder_chat_onboarding_shown=1; frontend-offer-card.newbuilding_broker_onboarding_shown=1; frontend-serp.chatAnimationCounter=10; frontend-serp.chatAnimationPrevPath=%2Fkupit-kvartiru-novostroyki%2F; countCallNowPopupShowed=1%3A1775476258941; forever_region_id=4959; forever_region_name=%D0%A0%D0%BE%D1%81%D1%82%D0%BE%D0%B2%D1%81%D0%BA%D0%B0%D1%8F%20%D0%BE%D0%B1%D0%BB%D0%B0%D1%81%D1%82%D1%8C; forever_main_town_region_id=4959; session_region_id=4593; session_main_town_region_id=1; login_mro_popup=1; _ym_isad=2; newbuilding_mortgage_payment_filter_onboarding=1; sopr_session=984fe57857c04e92; _ym_visorc=b; _yasc=TRKPtQwX100Llb+e03dywP4BMM51qmiqnpIri2vjwENZelIZncfgiKhy9oA6lHkCD45B0dE=; _yasc=0SMjQFnVkpPLd+wWAfEiWFOgNTsWxvn3eJZjAgkulSdK+kvhtbHsGVU0Gbe2z9s/OyK/5r8=; _ga_3369S417EL=GS2.1.s1775561722$o34$g1$t1775562653$j60$l0$h0',
    }

    json_data = {
        'jsonQuery': {
            '_type': 'flatsale',
            'engine_version': {
                'type': 'term',
                'value': 2,
            },
            'sort': {
                'type': 'term',
                'value': 'price_object_order',
            },
            'geo': {
                'type': 'geo',
                'value': [
                    {
                        'type': 'newobject',
                        'id': 4825183,
                    },
                ],
            },
            'decorations_list': {
                'type': 'terms',
                'value': [
                    'preFine',
                ],
            },
            'from_developer': {
                'type': 'term',
                'value': True,
            },
            'page': {
                'type': 'term',
                'value': 1,
            },
        },
    '_liquiditySource': 'web_serp',
    }



    def extract_digits_or_original(s):
        digits = ''.join([char for char in s if char.isdigit()])
        return int(digits) if digits else s

    current_date = datetime.date.today()

    no_flats = []

    for y in parsim:

        session = requests.Session()

        flats = []

        json_data["jsonQuery"]["page"]["value"] = 1

        print("Новый ЖК", y)

        json_data["jsonQuery"]["geo"]["value"][0]["id"] = y

        for decoration in decoration_list:

            counter = 1
            total_count = 1
            json_data["jsonQuery"]["decorations_list"]["value"][0] = decoration
            json_data["jsonQuery"]["page"]["value"] = 1
            print(decoration)
            print(json_data)




            while True:

                if counter > 1:
                    sleep_time = random.uniform(2, 4)
                    time.sleep(sleep_time)

                try:
                    response = session.post(
                        'https://api.cian.ru/search-offers/v2/search-offers-desktop/',
                        cookies=cookies,
                        headers=headers,
                        json=json_data,
                    )

                    print(response.status_code)
                except:
                    print("Произошла ошибка, пробуем ещё раз")
                    time.sleep(61)
                    session = requests.Session()
                    response = session.post(
                        'https://api.cian.ru/search-offers/v2/search-offers-desktop/',
                        cookies=cookies,
                        headers=headers,
                        json=json_data,
                    )
                    print(response.status_code)
                items = response.json()["data"]["offersSerialized"]


                for i in items:
                    try:
                        if i['building']['deadline']['isComplete'] == True:
                            srok_sdachi_old = "Дом сдан"
                        elif i['building']['deadline']['quarterEnd'] is None:
                            srok_sdachi_old = ''
                        else:
                            srok_sdachi_old = f"{i['building']['deadline']['quarterEnd']}"
                    except:
                        srok_sdachi_old = ''
                    try:
                        url = i['fullUrl']
                    except:
                        url = ''

                    try:
                        if i['isApartments'] == True:
                            type = "Апартаменты"
                        else:
                            type = "Квартира"
                    except:
                        type = ''
                    try:
                        if i['discount'] is not None:
                            price = extract_digits_or_original(i['discount']['newPrice'])
                            old_price = extract_digits_or_original(i['discount']['oldPrice'])
                        else:
                            old_price = i['bargainTerms']['priceRur']
                            price = ''
                    except:
                        old_price = i['bargainTerms']['priceRur']
                        price = ''
                    try:
                        project = i['geo']['jk']['displayName']
                    except:
                        project = ''
                    try:
                        finish_type = decoration_dict.get(decoration)
                    except:
                        finish_type = ''
                    try:
                        adress = i['geo']['userInput']
                    except:
                        adress = ""

                    try:
                        korpus = str(i["geo"]["jk"]["house"]["name"]).replace('Корпус ', '')
                    except:
                        korpus = ''

                    try:
                        developer = i['geo']['jk']['developer']['name']
                    except:
                        developer = ""

                    try:
                        if i["roomsCount"] == None:
                            room_count = 0
                        else:
                            room_count = int(i["roomsCount"])
                    except:
                        room_count = ''
                    try:
                        area = float(i["totalArea"])
                    except:
                        area = ''


                    date = datetime.date.today()

                    try:
                        floor = i["floorNumber"]
                    except:
                        floor = ''
                    english = ''
                    promzona = ''
                    mestopolozhenie = ''
                    subway = ''
                    distance_to_subway = ''
                    time_to_subway = ''
                    mck = ''
                    distance_to_mck = ''
                    time_to_mck = ''
                    bkl = ''
                    distance_to_bkl = ''
                    time_to_bkl = ''
                    status = ''
                    start = ''
                    comment = ''

                    okrug = ''
                    district = ''

                    eskrou = ''
                    konstruktiv = ''
                    klass = ''
                    srok_sdachi = ''
                    stadia = ''
                    dogovor = ''

                    discount = ''
                    price_per_metr = ''
                    price_per_metr_new = ''

                    section = ''
                    flat_number = ''


                    print(
                        f"{project}, {url}, дата: {date}, кол-во комнат: {room_count}, площадь: {area}, цена: {price}, срок сдачи: {srok_sdachi_old}, корпус: {korpus}, этаж: {floor}, {finish_type} ")
                    result = [date, project, english, promzona, mestopolozhenie, subway, distance_to_subway, time_to_subway,
                              mck, distance_to_mck, time_to_mck, distance_to_bkl,
                              time_to_bkl, bkl, status, start, comment, developer, okrug, district, adress, eskrou, korpus,
                              konstruktiv, klass, srok_sdachi, srok_sdachi_old,
                              stadia, dogovor, type, finish_type, room_count, area, price_per_metr, old_price, discount,
                              price_per_metr_new, price, section, floor, flat_number]
                    flats.append(result)

                total_count = response.json()["data"]["offerCount"]
                downloaded = len(flats)
                print(f'ID ЖК: {y}. Отделка: {decoration}. Загружено {downloaded} предложений из {total_count}')
                print("-----------------------------------------------------------------------------")
                if not items:
                    break
                json_data["jsonQuery"]["page"]["value"] += 1



                counter += 1

        if len(flats) > 0:
            save_flats_to_excel(flats, project, developer, kvartirografia=True)
        else:
            no_flats.append(y)

    print(f'Пустые проекты: {no_flats}')
    return f"Парсинг завершен для ID: {cian_id}"

def start_parsing():
    cian_id = entry.get()

    if not cian_id:
        messagebox.showerror("Ошибка", "Введите ID")
        return

    try:
        result = parse_cian(cian_id)

        messagebox.showinfo(
            "Успех",
            result
        )

    except Exception as e:
        messagebox.showerror(
            "Ошибка",
            str(e)
        )


root = tk.Tk()
root.title("Парсер ЦИАН")
root.geometry("400x200")

label = tk.Label(root, text="Введите ID:")
label.pack(pady=10)

entry = tk.Entry(root, width=30)
entry.pack(pady=10)

start_button = tk.Button(
    root,
    text="Запуск",
    command=start_parsing
)
start_button.pack(pady=20)

root.mainloop()
